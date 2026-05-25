import pandas as pd
import numpy as np
import xml.etree.ElementTree as ET
import openpyxl # type: ignore
import openpyxl.styles # type: ignore
import random
from datetime import datetime, timedelta

# Asegurar reproducibilidad de datos sintéticos
np.random.seed(42)
random.seed(42)

print("--- GENERANDO ARCHIVOS LOCALES DEL PROYECTO ---")

# =========================================================================
# 1. GENERAR inventario.csv (Min. 500 filas, 10% nulos, 5% duplicados)
# =========================================================================
base_rows = 600
id_productos = list(range(1, base_rows + 1))

# Inyectar un 5% de IDs duplicados
num_dups = int(base_rows * 0.05)
dup_ids = random.choices(id_productos, k=num_dups)
all_ids = id_productos + dup_ids

# Generar stock e inyectar un 10% de valores nulos (NaN)
total_rows = len(all_ids)
stock = [float(random.randint(5, 150)) for _ in range(total_rows)]
num_nulls = int(total_rows * 0.10)
null_indices = random.sample(range(total_rows), num_nulls)
for idx in null_indices:
    stock[idx] = np.nan

df_inventario = pd.DataFrame({
    'id_producto': all_ids,
    'stock': stock
})
# Mezclar filas para simular la llegada de un archivo crudo "sucio"
df_inventario = df_inventario.sample(frac=1).reset_index(drop=True)
df_inventario.to_csv('inventario.csv', index=False)
print("[ÉXITO] 'inventario.csv' generado con nulos y duplicados simulados.")

# =========================================================================
# 2. GENERAR logs_servidor.txt (2,000 a 3,000 líneas para parsing RegEx)
# =========================================================================
num_logs = 2500
methods = ['GET', 'POST', 'PUT', 'DELETE']
urls = ['/home', '/productos', '/carrito', '/checkout', '/perfil', '/api/v1/status']
statuses = [200, 200, 200, 200, 304, 404, 500]  # Ponderado a operaciones exitosas
start_time = datetime(2026, 5, 1)

with open('logs_servidor.txt', 'w') as f:
    for _ in range(num_logs):
        current_time = start_time + timedelta(seconds=random.randint(1, 2000000))
        ts_str = current_time.strftime('%Y-%m-%d %H:%M:%S')
        method = random.choice(methods)
        url = random.choice(urls)
        status = random.choice(statuses)
        resp_time = random.randint(15, 1500) # Milisegundos de respuesta
        ip = f"192.168.1.{random.randint(1, 254)}"
        # Formato de línea semiestructurada para parsing clásico de logs
        log_line = f"[{ts_str}] IP={ip} {method} {url} STATUS={status} TIME={resp_time}ms\n"
        f.write(log_line)
print("[ÉXITO] 'logs_servidor.txt' generado con 2,500 líneas listas para parsing.")

# =========================================================================
# 3. GENERAR catalogos.xml (Definiciones estructuradas de categorías)
# =========================================================================
root = ET.Element("catalogos")
categorias = {
    "Electrónica": ["Smartphone", "Laptop", "Televisor", "Audífonos", "Cargador"],
    "Ropa": ["Camisa", "Pantalón", "Chaqueta", "Zapatos", "Vestido"],
    "Hogar": ["Sofá", "Mesa", "Lámpara", "Espejo", "Licuadora"]
}
for cat_name, subcats in categorias.items():
    cat_elem = ET.SubElement(root, "categoria", nombre=cat_name)
    for sub in subcats:
        sub_elem = ET.SubElement(cat_elem, "producto_tipo")
        sub_elem.text = sub

tree = ET.ElementTree(root)
ET.indent(tree, space="  ", level=0) # Indentación limpia del archivo XML
tree.write("catalogos.xml", encoding="utf-8", xml_declaration=True)
print("[ÉXITO] 'catalogos.xml' generado correctamente.")

# =========================================================================
# 4. GENERAR metas_anuales.xlsx (KPIs de negocio por región estilizados)
# =========================================================================
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Metas_KPI"
ws.views.sheetView[0].showGridLines = True # Mostrar líneas de cuadrícula siempre

headers = ["Región", "Meta_Ventas_MXN", "Año", "KPI_Eficiencia"]
ws.append(headers)

regiones = ["Norte", "Sur", "Centro", "Este", "Oeste"]
for reg in regiones:
    ws.append([
        reg,
        random.randint(1500000, 5000000),
        2026,
        round(random.uniform(0.80, 0.99), 2)
    ])

# Estilos visuales corporativos
header_fill = openpyxl.styles.PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_font = openpyxl.styles.Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
thin_border = openpyxl.styles.Border(
    left=openpyxl.styles.Side(style='thin', color='D9D9D9'), right=openpyxl.styles.Side(style='thin', color='D9D9D9'),
    top=openpyxl.styles.Side(style='thin', color='D9D9D9'), bottom=openpyxl.styles.Side(style='thin', color='D9D9D9')
)

for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = openpyxl.styles.Alignment(horizontal="center", vertical="center")

for row in ws.iter_rows(min_row=2, max_row=6, min_col=1, max_col=4):
    for cell in row:
        cell.border = thin_border
        cell.font = openpyxl.styles.Font(name="Segoe UI", size=10)
        if cell.column == 2:
            cell.number_format = '$#,##0'
            cell.alignment = openpyxl.styles.Alignment(horizontal="right")
        elif cell.column == 4:
            cell.number_format = '0%'
            cell.alignment = openpyxl.styles.Alignment(horizontal="right")
        elif cell.column == 3:
            cell.alignment = openpyxl.styles.Alignment(horizontal="center")

# Ajuste automático del ancho de columnas
for col in ws.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = col[0].column_letter
    ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

wb.save("metas_anuales.xlsx")
print("[ÉXITO] 'metas_anuales.xlsx' guardado con formatos profesionales.")
print("--- TODOS LOS ARCHIVOS FUERON CREADOS CON ÉXITO ---")